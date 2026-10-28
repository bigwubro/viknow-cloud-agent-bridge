"""Render original document pages for full-page visual inspection."""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from collections.abc import Iterator
from contextlib import contextmanager
from dataclasses import dataclass
from pathlib import Path

DOCUMENT_INSPECT_DPI = 144
MAX_DOCUMENT_INSPECT_PAGES = 8


def document_page_count(source_path: Path) -> int:
    """Read the number of physical PDF pages without rendering or model calls."""
    import fitz

    with fitz.open(source_path) as document:
        return document.page_count


@dataclass(frozen=True)
class RenderedDocumentPage:
    """A temporary PNG rendered from one source document page."""

    path: Path
    page_number: int
    source_text: str


@dataclass(frozen=True)
class SpreadsheetRenderContext:
    """Zero-based worksheet coordinates of the retrieved spreadsheet evidence."""

    sheet_index: int
    row_index: int
    column_start: int
    column_end: int


@contextmanager
def render_document_page(
    source_path: Path,
    *,
    page_number: int,
    dpi: int = DOCUMENT_INSPECT_DPI,
    spreadsheet_context: SpreadsheetRenderContext | None = None,
) -> Iterator[RenderedDocumentPage]:
    """Render one one-based document page and clean its temporary files afterwards."""
    if page_number < 1:
        raise ValueError("document page number must be at least 1")
    if dpi != DOCUMENT_INSPECT_DPI:
        raise ValueError(f"document inspection must render at {DOCUMENT_INSPECT_DPI} DPI")
    if not source_path.is_file():
        raise FileNotFoundError(f"document source does not exist: {source_path}")

    with tempfile.TemporaryDirectory(prefix="viknow-document-inspect-") as temporary_directory:
        temporary_path = Path(temporary_directory)
        pdf_path = _pdf_for_source(source_path, temporary_path, spreadsheet_context)
        page_path = temporary_path / f"page-{page_number}.png"
        source_text = _render_pdf_page(
            pdf_path, page_number=page_number, dpi=dpi, output_path=page_path
        )
        yield RenderedDocumentPage(path=page_path, page_number=page_number, source_text=source_text)


def _pdf_for_source(
    source_path: Path,
    temporary_directory: Path,
    spreadsheet_context: SpreadsheetRenderContext | None,
) -> Path:
    if source_path.suffix.lower() == ".pdf":
        return source_path
    if source_path.suffix.lower() != ".xlsx":
        raise ValueError(
            "document inspection requires a PDF derivative or .xlsx source, "
            f"got {source_path.suffix}"
        )
    if spreadsheet_context is None:
        raise ValueError(
            ".xlsx document evidence requires a worksheet and retrieved cell or image location"
        )

    soffice = shutil.which("soffice")
    if soffice is None:
        raise RuntimeError("LibreOffice (soffice) is required to render .xlsx document evidence")
    xlsx_path = _prepare_spreadsheet_context(
        source_path,
        temporary_directory,
        spreadsheet_context,
    )
    completed = subprocess.run(
        [
            soffice,
            "--headless",
            "--convert-to",
            "pdf",
            "--outdir",
            str(temporary_directory),
            str(xlsx_path),
        ],
        check=False,
        capture_output=True,
        text=True,
        timeout=120,
    )
    pdf_path = temporary_directory / f"{xlsx_path.stem}.pdf"
    if completed.returncode != 0 or not pdf_path.is_file():
        detail = (completed.stderr or completed.stdout).strip()
        raise RuntimeError(
            f"failed to convert .xlsx source to PDF: {detail or completed.returncode}"
        )
    return pdf_path


def _prepare_spreadsheet_context(
    source_path: Path,
    temporary_directory: Path,
    context: SpreadsheetRenderContext,
) -> Path:
    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except ImportError as exc:  # pragma: no cover - supplied by HyperRAG dependency
        raise RuntimeError("openpyxl is required to render .xlsx document evidence") from exc

    workbook = load_workbook(source_path)
    if not 0 <= context.sheet_index < len(workbook.worksheets):
        raise ValueError(
            f"spreadsheet sheet {context.sheet_index} is unavailable; "
            f"workbook has {len(workbook.worksheets)} sheets"
        )
    worksheet = workbook.worksheets[context.sheet_index]
    for candidate in workbook.worksheets:
        candidate.sheet_state = "visible" if candidate is worksheet else "hidden"
    workbook.active = context.sheet_index

    if not worksheet.print_area:
        first_row = max(1, context.row_index + 1 - 20)
        last_row = context.row_index + 1 + 20
        first_column = max(1, context.column_start + 1 - 5)
        last_column = max(context.column_end, context.column_start + 1) + 5
        worksheet.print_area = (
            f"{get_column_letter(first_column)}{first_row}:"
            f"{get_column_letter(last_column)}{last_row}"
        )

    context_path = temporary_directory / f"{source_path.stem}-context.xlsx"
    workbook.save(context_path)
    return context_path


def _render_pdf_page(pdf_path: Path, *, page_number: int, dpi: int, output_path: Path) -> str:
    try:
        import fitz
    except ImportError as exc:  # pragma: no cover - dependency validation catches this
        raise RuntimeError("PyMuPDF is required to render document inspection pages") from exc

    with fitz.open(pdf_path) as document:
        if page_number > document.page_count:
            raise ValueError(
                f"document page {page_number} is unavailable; "
                f"document has {document.page_count} pages"
            )
        page = document[page_number - 1]
        page.get_pixmap(dpi=dpi).save(output_path)
        return page.get_text("text", sort=True)
