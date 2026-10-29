"""Extract reviewable information units from original source documents."""

import re
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

from .models import SourceInformationUnit


class SourceExtractionError(RuntimeError):
    pass


class SourceInformationExtractor:
    SUPPORTED_SUFFIXES = {".pdf", ".docx", ".pptx", ".xlsx", ".txt", ".md"}

    def extract(self, source_path: Path) -> list[SourceInformationUnit]:
        suffix = source_path.suffix.lower()
        if suffix not in self.SUPPORTED_SUFFIXES:
            raise SourceExtractionError(f"unsupported source format: {suffix}")
        if not source_path.is_file():
            raise SourceExtractionError(f"source file does not exist: {source_path}")
        if suffix in {".txt", ".md"}:
            return self._extract_text(source_path, markdown=suffix == ".md")
        if suffix == ".xlsx":
            return self._extract_xlsx(source_path)
        if suffix == ".docx":
            return self._extract_docx(source_path)
        if suffix == ".pptx":
            return self._extract_pptx(source_path)
        return self._extract_pdf(source_path)

    @staticmethod
    def _extract_text(path: Path, *, markdown: bool) -> list[SourceInformationUnit]:
        text = path.read_text(encoding="utf-8-sig")
        blocks = [block.strip() for block in re.split(r"\n\s*\n", text) if block.strip()]
        units: list[SourceInformationUnit] = []
        for index, block in enumerate(blocks):
            kind = "text"
            if markdown and re.match(r"^#{1,6}\s", block):
                kind = "heading"
            elif markdown and re.match(r"^(?:[-*+]\s|\d+[.)]\s)", block):
                kind = "list"
            elif markdown and "|" in block and re.search(r"\|\s*:?-{3,}", block):
                kind = "table"
            elif markdown and block.startswith("```"):
                kind = "code"
            units.append(
                SourceInformationUnit(
                    unit_id=f"block-{index}",
                    kind=kind,
                    locator={"block": index},
                    content=block,
                )
            )
        return units

    @staticmethod
    def _extract_xlsx(path: Path) -> list[SourceInformationUnit]:
        from openpyxl import load_workbook

        workbook = load_workbook(path, data_only=False, read_only=False)
        units: list[SourceInformationUnit] = []
        try:
            for sheet in workbook.worksheets:
                for row in sheet.iter_rows():
                    for cell in row:
                        if cell.value is None or str(cell.value).strip() == "":
                            continue
                        value = str(cell.value)
                        kind = "formula" if value.startswith("=") else "cell"
                        units.append(
                            SourceInformationUnit(
                                unit_id=f"{sheet.title}!{cell.coordinate}",
                                kind=kind,
                                locator={"sheet": sheet.title, "cell": cell.coordinate},
                                content=value,
                                weight=1.5 if kind == "formula" else 1.0,
                            )
                        )
                for index, merged_range in enumerate(sheet.merged_cells.ranges):
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"{sheet.title}!merged-{index}",
                            kind="merged_range",
                            locator={"sheet": sheet.title, "range": str(merged_range)},
                            content=str(merged_range),
                            reviewable=False,
                            risk_codes=["merged_cell_structure"],
                        )
                    )
                for index, _chart in enumerate(sheet._charts):
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"{sheet.title}!chart-{index}",
                            kind="chart",
                            locator={"sheet": sheet.title, "chart": index},
                            content="Embedded chart",
                            weight=2.0,
                            reviewable=False,
                            risk_codes=["visual_chart_requires_review"],
                        )
                    )
        finally:
            workbook.close()
        return units

    def _extract_docx(self, path: Path) -> list[SourceInformationUnit]:
        with zipfile.ZipFile(path) as archive:
            names = [
                name
                for name in archive.namelist()
                if name == "word/document.xml"
                or re.fullmatch(r"word/(?:header|footer)\d+\.xml", name)
            ]
            units: list[SourceInformationUnit] = []
            for name in names:
                root = ET.fromstring(archive.read(name))
                for index, paragraph in enumerate(root.iter(self._qn("w", "p"))):
                    text = "".join(
                        node.text or "" for node in paragraph.iter(self._qn("w", "t"))
                    ).strip()
                    if not text:
                        continue
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"{name}#p{index}",
                            kind="header_footer" if name != "word/document.xml" else "text",
                            locator={"part": name, "paragraph": index},
                            content=text,
                        )
                    )
                math_count = sum(1 for _ in root.iter(self._qn("m", "oMath")))
                for index in range(math_count):
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"{name}#math{index}",
                            kind="formula",
                            locator={"part": name, "formula": index},
                            content="Office Math formula",
                            weight=2.0,
                            reviewable=False,
                            risk_codes=["formula_conversion_loss"],
                        )
                    )
            return units

    def _extract_pptx(self, path: Path) -> list[SourceInformationUnit]:
        with zipfile.ZipFile(path) as archive:
            units: list[SourceInformationUnit] = []
            slide_names = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/slides/slide\d+\.xml", name)
            )
            for slide_index, name in enumerate(slide_names, start=1):
                root = ET.fromstring(archive.read(name))
                texts = [
                    (node.text or "").strip()
                    for node in root.iter(self._qn("a", "t"))
                    if (node.text or "").strip()
                ]
                if texts:
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"slide-{slide_index}-text",
                            kind="slide_text",
                            locator={"slide": slide_index},
                            content="\n".join(texts),
                            weight=max(1.0, len(texts) / 4),
                        )
                    )
                connector_count = sum(1 for _ in root.iter(self._qn("p", "cxnSp")))
                timing_count = sum(1 for _ in root.iter(self._qn("p", "timing")))
                visual_risks = ["visual_slide_requires_review"]
                if connector_count:
                    visual_risks.append("connector_semantics")
                if timing_count:
                    visual_risks.append("animation_semantics")
                units.append(
                    SourceInformationUnit(
                        unit_id=f"slide-{slide_index}-visual",
                        kind="visual_slide",
                        locator={"slide": slide_index},
                        content="PPT slide visual semantics",
                        weight=2.0,
                        reviewable=False,
                        risk_codes=visual_risks,
                    )
                )
            note_names = sorted(
                name
                for name in archive.namelist()
                if re.fullmatch(r"ppt/notesSlides/notesSlide\d+\.xml", name)
            )
            for note_index, name in enumerate(note_names, start=1):
                root = ET.fromstring(archive.read(name))
                note = " ".join(
                    (node.text or "").strip()
                    for node in root.iter(self._qn("a", "t"))
                    if (node.text or "").strip()
                )
                if note:
                    units.append(
                        SourceInformationUnit(
                            unit_id=f"slide-{note_index}-notes",
                            kind="speaker_notes",
                            locator={"slide": note_index},
                            content=note,
                        )
                    )
            return units

    @staticmethod
    def _extract_pdf(path: Path) -> list[SourceInformationUnit]:
        try:
            from pypdf import PdfReader
        except ImportError:
            from PyPDF2 import PdfReader

        reader = PdfReader(str(path))
        units: list[SourceInformationUnit] = []
        for page_index, page in enumerate(reader.pages):
            text = str(page.extract_text() or "").strip()
            if text:
                units.append(
                    SourceInformationUnit(
                        unit_id=f"page-{page_index}",
                        kind="page_text",
                        locator={"page": page_index},
                        content=text,
                        weight=max(1.0, len(text) / 1000),
                    )
                )
            units.append(
                SourceInformationUnit(
                    unit_id=f"page-{page_index}-visual",
                    kind="visual_page",
                    locator={"page": page_index},
                    content="PDF page visual semantics",
                    weight=2.0,
                    reviewable=False,
                    risk_codes=["visual_page_requires_review"],
                )
            )
        return units

    @staticmethod
    def _qn(prefix: str, local_name: str) -> str:
        namespaces = {
            "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
            "m": "http://schemas.openxmlformats.org/officeDocument/2006/math",
            "p": "http://schemas.openxmlformats.org/presentationml/2006/main",
            "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        }
        return f"{{{namespaces[prefix]}}}{local_name}"
